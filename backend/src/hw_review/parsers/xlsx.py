"""Read-only OOXML XLSX parser backed by openpyxl and ZIP inventory."""

from __future__ import annotations

import hashlib
import io
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from urllib.parse import quote
from uuid import NAMESPACE_URL, UUID, uuid5

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.exceptions import InvalidFileException

from hw_review.domain.hashing import normalized_text_hash, ordered_identity_hash
from hw_review.domain.models import (
    ContentBlock,
    DocumentContainer,
    ParseWarning,
    ReportDocument,
    StagedFile,
    TableCell,
)
from hw_review.parsers.base import ParserError


class XlsxParseError(ParserError):
    """Stable XLSX-specific parse failure."""


class XlsxParser:
    """Normalize a verified XLSX workbook without executing or recalculating it."""

    format = "XLSX"
    version = "openpyxl-3.1/read-only-contract-v1"

    def parse(self, staged: StagedFile) -> ReportDocument:
        if staged.detected_format != self.format:
            raise XlsxParseError("WRONG_FORMAT", "XLSX parser requires detected format XLSX")
        payload = self._verified_payload(staged)
        if not payload.startswith(b"PK"):
            raise XlsxParseError("CORRUPT_WORKBOOK", "XLSX payload is not an OOXML ZIP package")
        try:
            formulas = load_workbook(
                io.BytesIO(payload), data_only=False, read_only=False, keep_links=False
            )
            cached_values = load_workbook(
                io.BytesIO(payload), data_only=True, read_only=False, keep_links=False
            )
        except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
            raise XlsxParseError("CORRUPT_WORKBOOK", "XLSX workbook cannot be read") from error

        try:
            document_id = uuid5(NAMESPACE_URL, f"xlsx:{staged.id}:{staged.sha256}")
            containers, has_formula = self._sheet_containers(
                formulas, cached_values, document_id
            )
            objects, has_macro = self._inventory_package(payload)
            containers = self._attach_objects(containers, objects)
            warnings: list[ParseWarning] = []
            if has_formula:
                warnings.append(
                    ParseWarning(
                        code="FORMULA_NOT_RECALCULATED",
                        message=(
                            "formula text and any stored cached value were read; "
                            "no formula was executed or recalculated"
                        ),
                    )
                )
            if has_macro:
                warnings.append(
                    ParseWarning(
                        code="MACRO_CONTENT_NOT_EXECUTED",
                        message="macro content was inventoried but was not loaded or executed",
                    )
                )
            return ReportDocument(
                id=document_id,
                source_file_id=staged.id,
                format=self.format,
                parser_version=self.version,
                container_count=len(containers),
                text_digest=self._document_digest(containers),
                parse_warnings=tuple(warnings),
                containers=containers,
            )
        finally:
            formulas.close()
            cached_values.close()

    @staticmethod
    def _verified_payload(staged: StagedFile) -> bytes:
        try:
            payload = Path(staged.path).read_bytes()
        except OSError as error:
            raise XlsxParseError(
                "STAGED_FILE_CHANGED", "staged XLSX payload is unavailable"
            ) from error
        if len(payload) != staged.size_bytes:
            raise XlsxParseError("STAGED_FILE_CHANGED", "staged XLSX size has changed")
        if hashlib.sha256(payload).hexdigest() != staged.sha256:
            raise XlsxParseError("STAGED_FILE_CHANGED", "staged XLSX hash has changed")
        return payload

    def _sheet_containers(
        self, formulas, cached_values, document_id: UUID
    ) -> tuple[tuple[DocumentContainer, ...], bool]:
        containers: list[DocumentContainer] = []
        has_formula = False
        for sheet_index, sheet in enumerate(formulas.worksheets):
            cached_sheet = cached_values[sheet.title]
            escaped_name = quote(sheet.title, safe="")
            prefix = f"sheet:{sheet_index}:{escaped_name}"
            merged_by_master = {
                (item.min_row - 1, item.min_col - 1): str(item)
                for item in sheet.merged_cells.ranges
            }
            cells: list[TableCell] = []
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell, MergedCell) or cell.value is None:
                        continue
                    is_formula = cell.data_type == "f"
                    has_formula = has_formula or is_formula
                    raw_value = self._json_scalar(cell.value)
                    cached_value = (
                        self._json_scalar(cached_sheet[cell.coordinate].value)
                        if is_formula
                        else None
                    )
                    display_source = cached_value if cached_value is not None else raw_value
                    display_value = self._display_value(display_source)
                    structural_address = f"{prefix}/cell:{cell.coordinate}"
                    cells.append(
                        TableCell(
                            row=cell.row - 1,
                            column=cell.column - 1,
                            address=cell.coordinate,
                            structural_address=structural_address,
                            raw_value=raw_value,
                            display_value=display_value,
                            formula_if_available=str(cell.value) if is_formula else None,
                            cached_formula_value=cached_value,
                            merged_range=merged_by_master.get(
                                (cell.row - 1, cell.column - 1)
                            ),
                            content_hash=normalized_text_hash(display_value),
                        )
                    )
            blocks: tuple[ContentBlock, ...] = ()
            container_id = uuid5(document_id, f"container:{sheet_index}:{sheet.title}")
            if cells:
                structural_address = f"{prefix}/range:{cells[0].address}:{cells[-1].address}"
                block_hash = ordered_identity_hash(
                    f"{cell.structural_address}\0{cell.content_hash}" for cell in cells
                )
                blocks = (
                    ContentBlock(
                        id=uuid5(container_id, structural_address),
                        kind="table",
                        order=0,
                        structural_address=structural_address,
                        content_hash=block_hash,
                        cells=tuple(cells),
                    ),
                )
            containers.append(
                DocumentContainer(
                    id=container_id,
                    kind="sheet",
                    name_or_number=sheet.title,
                    order=sheet_index,
                    blocks=blocks,
                )
            )
        return tuple(containers), has_formula

    @staticmethod
    def _json_scalar(value):
        if isinstance(value, (datetime, date, time)):
            return value.isoformat()
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        return str(value)

    @staticmethod
    def _display_value(value) -> str:
        if value is True:
            return "TRUE"
        if value is False:
            return "FALSE"
        return "" if value is None else str(value)

    @staticmethod
    def _inventory_package(
        payload: bytes,
    ) -> tuple[tuple[tuple[str, str, str], ...], bool]:
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as package:
                objects: list[tuple[str, str, str]] = []
                names = sorted(package.namelist(), key=str.casefold)
                has_macro = any(name.casefold().endswith("vbaproject.bin") for name in names)
                for name in names:
                    lowered = name.casefold()
                    if not (
                        lowered.startswith("xl/media/")
                        or lowered.startswith("xl/embeddings/")
                        or lowered.endswith("vbaproject.bin")
                    ):
                        continue
                    digest = hashlib.sha256()
                    with package.open(name) as stream:
                        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                            digest.update(chunk)
                    kind = "image" if lowered.startswith("xl/media/") else "attachment"
                    objects.append((kind, name, digest.hexdigest()))
                return tuple(objects), has_macro
        except (zipfile.BadZipFile, KeyError, OSError) as error:
            raise XlsxParseError(
                "CORRUPT_WORKBOOK", "XLSX package inventory cannot be read"
            ) from error

    @staticmethod
    def _attach_objects(
        containers: tuple[DocumentContainer, ...],
        objects: tuple[tuple[str, str, str], ...],
    ) -> tuple[DocumentContainer, ...]:
        if not containers or not objects:
            return containers
        first = containers[0]
        blocks = list(first.blocks)
        for kind, name, content_hash in objects:
            address = f"workbook/object:{quote(name, safe='')}"
            blocks.append(
                ContentBlock(
                    id=uuid5(first.id, address),
                    kind=kind,
                    order=len(blocks),
                    structural_address=address,
                    content_hash=content_hash,
                )
            )
        return (first.model_copy(update={"blocks": tuple(blocks)}), *containers[1:])

    @staticmethod
    def _document_digest(containers: tuple[DocumentContainer, ...]) -> str:
        return ordered_identity_hash(
            f"{container.order}\0{block.order}\0{block.structural_address}\0{block.content_hash}"
            for container in containers
            for block in container.blocks
        )

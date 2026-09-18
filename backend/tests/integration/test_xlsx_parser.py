"""Integration coverage for the read-only OOXML XLSX adapter."""

from __future__ import annotations

import hashlib
from datetime import datetime
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook

from hw_review.domain.enums import FileRole
from hw_review.domain.models import StagedFile
from hw_review.parsers.xlsx import XlsxParseError, XlsxParser


def _write_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet 1 中文"
    sheet.merge_cells("B2:C2")
    sheet["B2"] = "结论"
    sheet["A3"] = 12.5
    sheet["A4"] = True
    sheet["A5"] = datetime(2026, 9, 17, 8, 30)
    sheet["A6"] = "=A3*2"
    workbook.create_sheet("Empty")
    workbook.save(path)
    workbook.close()


def _staged_for(path: Path, *, detected_format: str = "XLSX") -> StagedFile:
    payload = path.read_bytes()
    return StagedFile(
        id=uuid4(),
        task_id=uuid4(),
        role=FileRole.PRIMARY_REPORT,
        original_name=path.name,
        detected_format=detected_format,
        path=path,
        size_bytes=len(payload),
        sha256=hashlib.sha256(payload).hexdigest(),
        source_mtime_ns=path.stat().st_mtime_ns,
    )


def test_xlsx_preserves_sheet_cells_merges_formulas_and_digest(tmp_path: Path) -> None:
    path = tmp_path / "fixture.xlsx"
    _write_fixture(path)
    staged = _staged_for(path)
    parser = XlsxParser()

    first = parser.parse(staged)
    second = parser.parse(staged)

    assert [(item.name_or_number, item.order) for item in first.containers] == [
        ("Sheet 1 中文", 0),
        ("Empty", 1),
    ]
    conclusion = first.find_cell("Sheet 1 中文", "B2")
    assert conclusion.display_value == "结论"
    assert conclusion.merged_range == "B2:C2"
    assert conclusion.structural_address == "sheet:0:Sheet%201%20%E4%B8%AD%E6%96%87/cell:B2"
    assert first.find_cell("Sheet 1 中文", "A4").raw_value is True
    assert first.find_cell("Sheet 1 中文", "A5").raw_value == "2026-09-17T08:30:00"
    formula = first.find_cell("Sheet 1 中文", "A6")
    assert formula.formula_if_available == "=A3*2"
    assert formula.cached_formula_value is None
    assert formula.display_value == "=A3*2"
    assert any(item.code == "FORMULA_NOT_RECALCULATED" for item in first.parse_warnings)
    assert first.id == second.id
    assert first.text_digest == second.text_digest


def test_xlsx_inventory_records_media_without_extracting_it(tmp_path: Path) -> None:
    path = tmp_path / "media.xlsx"
    _write_fixture(path)
    # A workbook ZIP member is sufficient to prove deterministic, non-executing inventory.
    import zipfile

    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("xl/media/image1.png", b"not-decoded-image")

    document = XlsxParser().parse(_staged_for(path))
    objects = [
        block
        for container in document.containers
        for block in container.blocks
        if block.structural_address.startswith("workbook/object:")
    ]

    assert [(item.kind, item.structural_address) for item in objects] == [
        ("image", "workbook/object:xl%2Fmedia%2Fimage1.png")
    ]


@pytest.mark.parametrize(
    ("mutation", "code"),
    [("wrong-format", "WRONG_FORMAT"), ("changed", "STAGED_FILE_CHANGED"), ("corrupt", "CORRUPT_WORKBOOK")],
)
def test_xlsx_rejects_wrong_changed_or_corrupt_sources(
    tmp_path: Path, mutation: str, code: str
) -> None:
    path = tmp_path / "invalid.xlsx"
    _write_fixture(path)
    staged = _staged_for(path)
    if mutation == "wrong-format":
        staged = staged.model_copy(update={"detected_format": "XLS"})
    elif mutation == "changed":
        path.write_bytes(path.read_bytes() + b"changed")
    else:
        path.write_bytes(b"PK\x03\x04not-a-workbook")
        staged = _staged_for(path)

    with pytest.raises(XlsxParseError) as error:
        XlsxParser().parse(staged)

    assert error.value.code == code
